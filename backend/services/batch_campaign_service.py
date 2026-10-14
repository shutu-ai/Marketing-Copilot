import os
import json
import math
import re
import hashlib
import csv
import errno
import logging
import threading
from datetime import datetime
from database import read_transaction, write_transaction
from models.batch_campaign import BatchCampaignCreate, BatchCampaignUpdate
from typing import Optional
import shutil
import pandas as pd
import time
import uuid
from services import settings_service
from services import (sms_task_service, operations_service, campaign_name_service,
                      ingestion_job_service, backpressure_service, metrics_service)

logger = logging.getLogger(__name__)
OVERDUE_CANCEL_BATCH_SIZE = max(
    50, min(2000, int(os.environ.get("MC_BATCH_CANCEL_BATCH_SIZE", "250")))
)
OVERDUE_CANCEL_CAMPAIGNS_PER_RUN = max(
    1, min(32, int(os.environ.get("MC_BATCH_CANCEL_CAMPAIGNS_PER_RUN", "4")))
)
_overdue_campaign_cursor = 0
_batch_generation_lock = threading.RLock()
_batch_source_lock = threading.RLock()
_batch_source_cursors = {}
_batch_generation_cursors = {}

try:
    _BATCH_INGEST_QUANTUM = max(
        1, min(100, int(os.environ.get("MC_BATCH_INGEST_QUANTUM", "1")))
    )
except (TypeError, ValueError):
    _BATCH_INGEST_QUANTUM = 1

UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "..", "uploads", "campaigns")
ALLOWED_LANGUAGES = {"amharic", "tigrigna", "somali", "english", "oromign", "en"}
BATCH_CUSTOMER_LIST_EXTENSIONS = {".csv", ".txt", ".xlsx", ".xls"}
CUSTOMER_LIST_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "latin-1")
INGEST_QUEUE_HIGH_WATERMARK = backpressure_service.TASK_HIGH
INGEST_QUEUE_LOW_WATERMARK = backpressure_service.TASK_LOW


def _run_campaign_write(operation: str, writer):
    return write_transaction(
        writer,
        attempts=3,
        operation=f"batch_campaign.{operation}",
        base_delay=0.1,
    )


def _queued_task_count() -> int:
    return backpressure_service._task_depth()


def _wait_for_ingest_capacity(stop_event=None) -> bool:
    """Pause on task, buffered-write or disk pressure with low-water recovery."""
    queued = _queued_task_count()
    if queued >= INGEST_QUEUE_HIGH_WATERMARK:
        while queued > INGEST_QUEUE_LOW_WATERMARK:
            if stop_event is not None and stop_event.is_set():
                return False
            time.sleep(0.25)
            queued = _queued_task_count()
    return backpressure_service.wait_for_ingestion(stop_event=stop_event)


class _CustomerListReadStats:
    def __init__(self):
        self.malformed_rows = 0
        self.invalid_msisdn_rows = 0

    @property
    def ignored_rows(self) -> int:
        return self.malformed_rows + self.invalid_msisdn_rows

    def bad_csv_line(self, _bad_line):
        self.malformed_rows += 1
        return None


def _read_delimited_sample(path: str) -> tuple[str, str]:
    last_error = None
    for encoding in CUSTOMER_LIST_ENCODINGS:
        try:
            with open(path, "r", encoding=encoding, newline="") as handle:
                return handle.read(8192), encoding
        except UnicodeDecodeError as exc:
            last_error = exc
    if last_error:
        raise ValueError("Unable to decode customer list file") from last_error
    return "", "utf-8-sig"


def _detect_csv_format(path: str) -> tuple[str, str]:
    sample, encoding = _read_delimited_sample(path)
    lines = [line for line in sample.splitlines() if line.strip()]
    first_line = lines[0].strip() if lines else ""
    counts = {delimiter: first_line.count(delimiter) for delimiter in (",", "|", ";", "\t")}
    likely = max(counts, key=counts.get)
    if counts[likely] > 0:
        return likely, encoding
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",|;\t").delimiter
    except csv.Error:
        delimiter = ","
    return delimiter, encoding


def _detect_csv_delimiter(path: str) -> str:
    try:
        return _detect_csv_format(path)[0]
    except Exception:
        return ","


def _canonicalize_customer_list_columns(df: pd.DataFrame) -> pd.DataFrame:
    """Normalize the reserved customer-list headers without changing order.

    ``msisdn`` and ``language`` are logical field names, not case-sensitive
    source headers.  Keep other columns intact because they may be used as
    template variables, but always expose the two reserved fields under their
    canonical lower-case names to the validation and enqueue paths.
    """
    rename = {}
    reserved = {}
    for column in df.columns:
        key = str(column).lstrip("\ufeff").strip().casefold()
        if key not in {"msisdn", "language"}:
            continue
        if key in reserved:
            raise ValueError(
                f"Customer list file contains duplicate {key} columns"
            )
        reserved[key] = column
        rename[column] = key
    return df.rename(columns=rename)


def _read_excel_chunks(
    path: str, chunksize: int, stats: Optional[_CustomerListReadStats] = None,
):
    """Read Excel input in bounded chunks and close workbook resources."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        # pandas delegates legacy XLS files to xlrd.  Keep the dependency
        # explicit so offline deployments fail with a useful message if an
        # old wheelhouse has not yet been refreshed.
        try:
            frame = pd.read_excel(path, dtype=str, keep_default_na=False, engine="xlrd")
        except ImportError as exc:
            raise ValueError("XLS files require the xlrd dependency") from exc
        for start in range(0, len(frame), max(1, chunksize)):
            yield frame.iloc[start:start + max(1, chunksize)].copy()
        return

    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if not header_row:
            return
        headers = [str(value or "").strip() for value in header_row]
        buffer = []
        for values in rows:
            values = tuple(values)
            if len(values) != len(headers):
                # openpyxl normally pads rows to the sheet width, but retain
                # this guard for malformed workbooks.
                if stats is not None:
                    stats.malformed_rows += 1
                continue
            buffer.append(dict(zip(headers, values)))
            if len(buffer) >= chunksize:
                yield pd.DataFrame(buffer, columns=headers)
                buffer.clear()
        if buffer:
            yield pd.DataFrame(buffer, columns=headers)
    finally:
        workbook.close()


def _prepare_customer_list_frame(
    df: pd.DataFrame,
    stats: Optional[_CustomerListReadStats] = None,
    row_offset: int = 0,
) -> pd.DataFrame:
    """Canonicalize a chunk and drop only rows with unusable MSISDNs.

    Structural errors (missing/duplicate reserved columns) and invalid
    language values remain hard failures.  An empty or unparsable MSISDN is a
    row-level data problem and is deliberately excluded from the returned
    frame so validation and later batch generation see the same rows.
    """
    stats = stats or _CustomerListReadStats()
    df = _canonicalize_customer_list_columns(df)
    columns = set(df.columns)
    if "msisdn" not in columns:
        raise ValueError("Customer list file must contain an msisdn column")
    if df.empty:
        return df.copy()

    valid_positions = []
    normalized_values = []
    for idx, value in enumerate(df["msisdn"]):
        try:
            normalized = sms_task_service.normalize_msisdn(value)
        except Exception:
            normalized = ""
        if not normalized:
            stats.invalid_msisdn_rows += 1
            continue
        valid_positions.append(idx)
        normalized_values.append(normalized)

    if not valid_positions:
        return df.iloc[0:0].copy()

    prepared = df.iloc[valid_positions].copy()
    prepared["msisdn"] = normalized_values
    if "language" in columns:
        allowed_label = ", ".join(sorted(ALLOWED_LANGUAGES))
        for idx, value in enumerate(prepared["language"]):
            language = str(value).strip().lower() if pd.notna(value) else ""
            if not language or language not in ALLOWED_LANGUAGES:
                source_row = row_offset + valid_positions[idx] + 2
                raise ValueError(
                    f"Row {source_row}: language '{str(value).strip()}' is invalid; "
                    f"allowed: {allowed_label}"
                )
    return prepared


def _iter_customer_list_frames(
    path: str, chunksize: int = 10000, stats: Optional[_CustomerListReadStats] = None,
):
    ext = os.path.splitext(path)[1].lower()
    if ext not in BATCH_CUSTOMER_LIST_EXTENSIONS:
        raise ValueError("Only CSV, TXT, XLSX and XLS files are supported")
    stats = stats or _CustomerListReadStats()
    row_offset = 0
    if ext in {".csv", ".txt"}:
        delimiter, encoding = _detect_csv_format(path)
        source = pd.read_csv(
            path, dtype=str, sep=delimiter, encoding=encoding,
            chunksize=max(1, chunksize), keep_default_na=False,
            engine="python", on_bad_lines=stats.bad_csv_line,
        )
    else:
        source = _read_excel_chunks(path, max(1, chunksize), stats)
    for frame in source:
        prepared = _prepare_customer_list_frame(frame, stats, row_offset)
        row_offset += len(frame)
        if not prepared.empty:
            yield prepared
        del frame, prepared


def _validate_customer_list_file(path: str, chunksize: int = 10000) -> dict:
    stats = _CustomerListReadStats()
    total = 0
    for frame in _iter_customer_list_frames(path, chunksize, stats):
        total += len(frame)
    if total <= 0:
        raise ValueError("Customer list file contains no valid MSISDN rows")
    return {
        "row_count": total,
        "ignored_rows": stats.ignored_rows,
        "ignored_msisdn_rows": stats.invalid_msisdn_rows,
        "ignored_malformed_rows": stats.malformed_rows,
    }


def _get_batch_base_dir() -> str:
    return settings_service.get_setting("batch_base_dir") or os.path.join(
        os.path.dirname(__file__), "..", "uploads", "batch_campaigns"
    )


def _sanitize_dirname(name: str) -> str:
    s = re.sub(r'[<>:"/\\|?*]', "_", name).strip()
    s = s.rstrip(". ")
    if not s:
        s = "campaign"
    return s


def _ensure_campaign_dir(campaign: dict) -> str:
    base = _get_batch_base_dir()
    dirname = _sanitize_dirname(campaign.get("name", f"campaign_{campaign['id']}"))
    campaign_dir = os.path.join(base, dirname)
    return campaign_dir


def _fmt_days(days: Optional[list]) -> str:
    return json.dumps(days, ensure_ascii=False) if days else "[]"


def _parse_days(val) -> list:
    if not val:
        return []
    if isinstance(val, list):
        return val
    if isinstance(val, str):
        try:
            return json.loads(val)
        except Exception:
            return []
    return []


def _get_campaign_templates(batch_campaign_id: int) -> list:
    rows = read_transaction(
        lambda conn: conn.execute("""
            SELECT st.id, st.language, st.content, st.source_addr
            FROM sms_templates st
            JOIN campaigns c ON c.id = st.campaign_id
            WHERE c.marketing_campaign_id = ? AND c.marketing_type = 'batch'
            ORDER BY st.language
        """, (batch_campaign_id,)).fetchall(),
        operation="batch_campaign.templates.list",
    )
    return [dict(r) for r in rows]


def _enrich(c: dict) -> dict:
    c["send_days"] = _parse_days(c.get("send_days", ""))
    return c


def get_all() -> list:
    rows = read_transaction(
        lambda conn: conn.execute(
            "SELECT * FROM batch_campaigns ORDER BY created_at DESC"
        ).fetchall(),
        operation="batch_campaign.list",
    )
    return [_enrich(dict(r)) for r in rows]


def list_campaigns(page: int = 1, page_size: int = 50) -> dict:
    def read(conn):
        total = conn.execute("SELECT COUNT(*) cnt FROM batch_campaigns").fetchone()["cnt"]
        offset = max(0, page - 1) * page_size
        rows = conn.execute(
            "SELECT * FROM batch_campaigns ORDER BY created_at DESC LIMIT ? OFFSET ?",
            (page_size, offset),
        ).fetchall()
        return total, rows

    total, rows = read_transaction(read, operation="batch_campaign.list_page")
    return {
        "total": total,
        "page": page,
        "page_size": page_size,
        "items": [_enrich(dict(r)) for r in rows],
    }


def get_one(campaign_id: int) -> Optional[dict]:
    row = read_transaction(
        lambda conn: conn.execute(
            "SELECT * FROM batch_campaigns WHERE id=?", (campaign_id,)
        ).fetchone(),
        operation="batch_campaign.get",
    )
    return _enrich(dict(row)) if row else None


def create(data: BatchCampaignCreate) -> dict:
    name = campaign_name_service.ensure_unique_name(data.name)
    if not (data.phone_prefix or "").strip():
        raise ValueError("请输入号码前缀")
    use_custom_directory = bool(data.use_custom_directory and (data.campaign_dir or "").strip())
    campaign_dir = (data.campaign_dir or "").strip() if use_custom_directory else ""

    def write(conn):
        cursor = conn.execute(
            "INSERT INTO batch_campaigns (name, start_time, end_time, send_start_time, send_end_time, send_days, batch_size, campaign_dir, use_custom_directory, sms_template_id, smpp_config_id, flash_sms, phone_prefix) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (name, data.start_time or "", data.end_time or "",
             data.send_start_time or "", data.send_end_time or "",
             _fmt_days(data.send_days), data.batch_size or 500, campaign_dir,
             use_custom_directory, data.sms_template_id, data.smpp_config_id,
             bool(data.flash_sms), data.phone_prefix or ""),
        )
        return dict(conn.execute(
            "SELECT * FROM batch_campaigns WHERE id=?", (cursor.lastrowid,)
        ).fetchone())

    campaign = _enrich(_run_campaign_write("create", write))
    if not campaign_dir:
        campaign_dir = _ensure_campaign_dir(campaign)
        _run_campaign_write(
            "create.directory",
            lambda conn: conn.execute(
                "UPDATE batch_campaigns SET campaign_dir=? WHERE id=?",
                (campaign_dir, campaign["id"]),
            ),
        )
        campaign["campaign_dir"] = campaign_dir
    os.makedirs(os.path.join(campaign_dir, "whitelist"), exist_ok=True)
    os.makedirs(os.path.join(campaign_dir, "output"), exist_ok=True)
    return campaign


def update(campaign_id: int, data: BatchCampaignUpdate) -> Optional[dict]:
    existing = read_transaction(
        lambda conn: conn.execute(
            "SELECT * FROM batch_campaigns WHERE id=?", (campaign_id,)
        ).fetchone(),
        operation="batch_campaign.update_probe",
    )
    if not existing:
        return None

    existing_dict = dict(existing)
    if data.name is not None and data.name.strip() != (existing_dict.get("name") or "").strip():
        data.name = campaign_name_service.ensure_unique_name(data.name, "batch", campaign_id)
    auto_campaign_dir = ""
    if data.status == "in_progress":
        validation_campaign = dict(existing_dict)
        if not (validation_campaign.get("campaign_dir") or "").strip() and not bool(validation_campaign.get("use_custom_directory")):
            auto_campaign_dir = _ensure_campaign_dir(validation_campaign)
            validation_campaign["campaign_dir"] = auto_campaign_dir
        validation_errors = operations_service.validate_campaign(
            "batch", {**validation_campaign, "id": campaign_id}
        )
        if validation_errors:
            raise ValueError("; ".join(validation_errors))
        if not existing_dict.get("row_count", 0):
            raise ValueError("请先上传潜客清单")
        if not existing_dict.get("sms_template_id"):
            raise ValueError("请先配置短信模板")

    if data.use_custom_directory is not None and existing_dict.get("customer_list_path"):
        new_dir = ((data.campaign_dir or "").strip()
                   if data.use_custom_directory
                   else _ensure_campaign_dir(existing_dict))
        old_dir = (existing_dict.get("campaign_dir") or "").strip()
        if new_dir and os.path.normcase(os.path.normpath(new_dir)) != os.path.normcase(os.path.normpath(old_dir)):
            raise ValueError("Cannot change campaign directory after a target list has been uploaded")
    if existing_dict.get("status") == "ended":
        only_allowed = all(
            getattr(data, f, None) is None
            for f in ["start_time", "batch_size", "campaign_dir", "use_custom_directory"]
        )
        if not only_allowed:
            raise ValueError("Ended campaign can only edit name, end time, send window, send days, and SMS template")
    elif existing_dict.get("status") != "not_started":
        only_allowed = all(
            getattr(data, f, None) is None
            for f in ["name", "start_time", "end_time", "send_start_time", "send_end_time", "send_days", "batch_size", "campaign_dir", "use_custom_directory", "flash_sms"]
        )
        if not only_allowed:
            raise ValueError("Campaign has already started - cannot edit")

    fields = {}
    for field in ["name", "start_time", "end_time", "status", "send_start_time", "send_end_time", "batch_size", "campaign_dir", "use_custom_directory", "sms_template_id", "smpp_config_id", "flash_sms", "phone_prefix"]:
        value = getattr(data, field, None)
        if value is not None:
            fields[field] = value
    if data.send_days is not None:
        fields["send_days"] = _fmt_days(data.send_days)
    if auto_campaign_dir:
        fields["campaign_dir"] = auto_campaign_dir
    if data.use_custom_directory is not None:
        if data.use_custom_directory:
            campaign_dir = (data.campaign_dir or "").strip()
            if not campaign_dir:
                raise ValueError("Custom campaign directory is required")
            fields["campaign_dir"] = campaign_dir
        else:
            fields["campaign_dir"] = _ensure_campaign_dir({**existing_dict, **fields})
    new_smpp_id = fields.get("smpp_config_id")
    status_update = "status" in fields and fields["status"] != existing_dict.get("status")

    def write(conn):
        current = conn.execute(
            "SELECT * FROM batch_campaigns WHERE id=?", (campaign_id,)
        ).fetchone()
        if not current:
            return None
        if new_smpp_id is not None:
            smpp = conn.execute(
                "SELECT id FROM smpp_configs WHERE id=? AND enabled=1", (new_smpp_id,)
            ).fetchone()
            if not smpp:
                raise ValueError("Associated SMPP configuration is missing or disabled")
        current_dict = dict(current)
        status_changed = "status" in fields and fields["status"] != current_dict.get("status")
        next_generation = int(current_dict.get("dispatch_generation") or 0)
        if status_changed:
            next_generation += 1
        if fields:
            write_fields = dict(fields)
            if status_changed:
                write_fields["dispatch_generation"] = next_generation
            set_clause = ", ".join(f"{key}=?" for key in write_fields)
            conn.execute(
                f"UPDATE batch_campaigns SET {set_clause}, updated_at=datetime('now','localtime') WHERE id=?",
                (*write_fields.values(), campaign_id),
            )
        if status_changed:
            sms_task_service.reconcile_campaign_state_in_transaction(
                conn, "batch", campaign_id, fields["status"], next_generation,
            )
        if new_smpp_id is not None:
            conn.execute(
                """UPDATE sms_tasks SET smpp_config_id=?, updated_at=CURRENT_TIMESTAMP
                   WHERE campaign_type='batch' AND campaign_id=? AND status IN ('queued','retrying','sending')""",
                (new_smpp_id, campaign_id),
            )
        row = conn.execute(
            "SELECT * FROM batch_campaigns WHERE id=?", (campaign_id,)
        ).fetchone()
        return dict(row) if row else None

    from contextlib import nullcontext
    from services import fast_send_service
    barrier = (
        fast_send_service.campaign_control_barrier("batch", campaign_id)
        if status_update and fast_send_service.is_fast() else nullcontext()
    )
    with barrier:
        updated_campaign = _run_campaign_write("update", write)
        if status_update and fast_send_service.is_fast() and updated_campaign:
            fast_send_service.publish_campaign_state(
                "batch", campaign_id, updated_campaign.get("status", ""),
                int(updated_campaign.get("dispatch_generation") or 0),
            )
    if updated_campaign is None:
        return None
    if fields.get("status") is not None:
        try:
            from services.smpp_gateway import gateway
            gateway.invalidate_campaign(
                "batch", campaign_id,
                int(updated_campaign.get("dispatch_generation") or 0),
            )
        except Exception:
            logger.debug("Unable to invalidate live batch SMPP queues", exc_info=True)
    sms_task_service.invalidate_campaign_cache("batch", campaign_id)
    updated_campaign = _enrich(updated_campaign)
    if data.use_custom_directory is not None:
        os.makedirs(os.path.join(updated_campaign["campaign_dir"], "whitelist"), exist_ok=True)
        os.makedirs(os.path.join(updated_campaign["campaign_dir"], "output"), exist_ok=True)
    if data.status == "in_progress":
        generate_batches(campaign_id, updated_campaign.get("batch_size") or 500)
    return updated_campaign


def get_prospects(campaign_id: int, search: str = "", field: str = "", value: str = "", page: int = 1, page_size: int = 50) -> dict:
    campaign = get_one(campaign_id)
    if not campaign or not campaign.get("customer_list_path"):
        return {"headers": [], "rows": [], "total": 0, "page": 1, "page_size": page_size}

    path = campaign["customer_list_path"]
    if not os.path.exists(path):
        return {"headers": [], "rows": [], "total": 0, "page": 1, "page_size": page_size}

    try:
        frames = list(_iter_customer_list_frames(path, chunksize=10000))
        df = pd.concat(frames, ignore_index=True).fillna("") if frames else pd.DataFrame()
    except Exception:
        return {"headers": [], "rows": [], "total": 0, "page": 1, "page_size": page_size}

    # Apply filters
    if search:
        mask = df.apply(lambda row: row.astype(str).str.contains(search, case=False, na=False).any(), axis=1)
        df = df[mask]

    if field and value:
        if field in df.columns:
            df = df[df[field].astype(str).str.contains(value, case=False, na=False)]

    total = len(df)
    start = (page - 1) * page_size
    end = start + page_size
    page_data = df.iloc[start:end].to_dict(orient="records")

    return {
        "headers": list(df.columns),
        "rows": page_data,
        "total": total,
        "page": page,
        "page_size": page_size,
    }


def delete(campaign_id: int) -> bool:
    row = read_transaction(
        lambda conn: conn.execute(
            "SELECT * FROM batch_campaigns WHERE id=?", (campaign_id,)
        ).fetchone(),
        operation="batch_campaign.delete_probe",
    )
    if not row:
        return False

    _run_campaign_write(
        "delete",
        lambda conn: (
            conn.execute("DELETE FROM batch_campaign_daily_stats WHERE campaign_id=?", (campaign_id,)),
            conn.execute("DELETE FROM batch_campaign_batches WHERE campaign_id=?", (campaign_id,)),
            conn.execute("DELETE FROM batch_campaigns WHERE id=?", (campaign_id,)),
        ),
    )
    campaign_dir = dict(row).get("campaign_dir", "")
    if campaign_dir and os.path.isdir(campaign_dir):
        try:
            shutil.rmtree(campaign_dir)
        except Exception:
            logger.warning("Unable to remove deleted batch campaign directory=%s", campaign_dir, exc_info=True)
    sms_task_service.invalidate_campaign_cache("batch", campaign_id)
    return True


def copy(campaign_id: int) -> Optional[dict]:
    def write(conn):
        row = conn.execute(
            "SELECT * FROM batch_campaigns WHERE id=?", (campaign_id,)
        ).fetchone()
        if not row:
            return None
        orig = dict(row)
        cursor = conn.execute(
            "INSERT INTO batch_campaigns (name, start_time, end_time, send_start_time, send_end_time, send_days, batch_size, sms_template_id, smpp_config_id, flash_sms, phone_prefix) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (orig["name"] + " (Copy)", orig["start_time"], orig["end_time"],
             orig.get("send_start_time", ""), orig.get("send_end_time", ""),
             orig.get("send_days", "[]"), orig.get("batch_size", 500),
             orig.get("sms_template_id"), orig.get("smpp_config_id"),
             bool(orig.get("flash_sms")), orig.get("phone_prefix", "")),
        )
        new_row = conn.execute(
            "SELECT * FROM batch_campaigns WHERE id=?", (cursor.lastrowid,)
        ).fetchone()
        return {"original": orig, "campaign": dict(new_row)}

    result = _run_campaign_write("copy", write)
    if result is None:
        return None
    orig = result["original"]
    new_campaign = _enrich(result["campaign"])
    campaign_dir = _ensure_campaign_dir(new_campaign)
    os.makedirs(os.path.join(campaign_dir, "whitelist"), exist_ok=True)
    os.makedirs(os.path.join(campaign_dir, "output"), exist_ok=True)
    old_dir = orig.get("campaign_dir") or _ensure_campaign_dir(orig)
    old_whitelist = os.path.join(old_dir, "whitelist")
    if os.path.isdir(old_whitelist):
        for fname in os.listdir(old_whitelist):
            src = os.path.join(old_whitelist, fname)
            if os.path.isfile(src):
                shutil.copy2(src, os.path.join(campaign_dir, "whitelist", fname))
    _run_campaign_write(
        "copy.directory",
        lambda conn: conn.execute(
            "UPDATE batch_campaigns SET campaign_dir=? WHERE id=?",
            (campaign_dir, new_campaign["id"]),
        ),
    )
    new_campaign["campaign_dir"] = campaign_dir
    return new_campaign


def get_daily_stats(campaign_id: int) -> list:
    rows = read_transaction(
        lambda conn: conn.execute(
            "SELECT * FROM batch_campaign_daily_stats WHERE campaign_id=? ORDER BY stats_date DESC",
            (campaign_id,)
        ).fetchall(),
        operation="batch_campaign.daily_stats",
    )
    return [dict(r) for r in rows]


def reset_batches_for_today(campaign_id: int) -> bool:
    """Mark all batches pending again for the current day (daily-repeat send).

    Only resets when this campaign has not enqueued anything today, so each
    eligible day produces exactly one batch run. Queued/retrying tasks left over
    from previous days are expired first so they cannot send late and duplicate
    the new day's batch.
    """
    today = datetime.now().strftime("%Y-%m-%d")
    def reset(conn):
        existing = conn.execute(
            "SELECT 1 FROM batch_campaign_daily_stats WHERE campaign_id=? AND stats_date=? LIMIT 1",
            (campaign_id, today),
        ).fetchone()
        if existing:
            return False
        conn.execute(
            """UPDATE sms_tasks SET status='expired', error_msg='superseded by next daily batch',
               updated_at=CURRENT_TIMESTAMP
               WHERE campaign_type='batch' AND campaign_id=? AND status IN ('queued','retrying')
                 AND CAST(created_at AS DATE) < ?""",
            (campaign_id, today),
        )
        count = conn.execute(
            """SELECT COUNT(*) cnt FROM batch_campaign_batches
               WHERE campaign_id=? AND status IN ('queued','sent','skipped','failed')""",
            (campaign_id,),
        ).fetchone()["cnt"]
        if not count:
            return False
        conn.execute(
            """UPDATE batch_campaign_batches
               SET status='pending',task_count=0,delivered_count=0,failed_count=0,
                   updated_at=CURRENT_TIMESTAMP WHERE campaign_id=?""",
            (campaign_id,),
        )
        conn.execute(
            """UPDATE ingestion_chunks SET status='pending',lease_token='',lease_expires_at=NULL,
               created_tasks=0,duplicate_rows=0,excluded_rows=0,invalid_rows=0,no_template_rows=0,
               error_detail='',completed_at=NULL,updated_at=CURRENT_TIMESTAMP
               WHERE source_batch_id IN
                 (SELECT id FROM batch_campaign_batches WHERE campaign_id=?)""",
            (campaign_id,),
        )
        conn.execute(
            """UPDATE ingestion_files SET status='pending',completed_chunks=0,created_tasks=0,
               error_detail='',completed_at=NULL,updated_at=CURRENT_TIMESTAMP
               WHERE campaign_type='batch' AND campaign_id=?""",
            (campaign_id,),
        )
        return True

    return bool(write_transaction(
        reset,
        attempts=3,
        operation="batch_campaign.reset_today",
    ))


def reset_daily_batches() -> int:
    """Reset batches of in-progress batch campaigns that have not run yet today."""
    today = datetime.now().strftime("%Y-%m-%d")
    rows = read_transaction(
        lambda conn: conn.execute(
            """SELECT c.id FROM batch_campaigns c
               WHERE c.status='in_progress' AND c.customer_list_path <> ''
                 AND EXISTS (SELECT 1 FROM batch_campaign_batches b WHERE b.campaign_id=c.id)
                 AND NOT EXISTS (SELECT 1 FROM batch_campaign_daily_stats d
                                 WHERE d.campaign_id=c.id AND d.stats_date=?)""",
            (today,),
        ).fetchall(),
        operation="batch_campaign.reset_candidates",
    )
    reset = 0
    for row in rows:
        if reset_batches_for_today(row["id"]):
            reset += 1
    if reset:
        logger.info("Batch daily reset prepared campaigns=%s", reset)
    return reset


def cancel_overdue_daily_tasks() -> int:
    """Cancel the day's still-unfinished batch tasks once the daily send window closes.

    Tasks already picked up by a worker ('sending') and submitted ones are left alone:
    they are already in flight at the SMSC. The next eligible day's batch run will
    re-send the same users.
    """
    now = datetime.now()
    current = now.strftime("%H:%M")
    global _overdue_campaign_cursor

    def load_candidates(conn, after_id):
        return conn.execute(
            """SELECT c.id, c.send_start_time, c.send_end_time,
                      EXISTS(
                          SELECT 1 FROM sms_tasks t
                          WHERE t.campaign_type='batch' AND t.campaign_id=c.id
                            AND t.status IN ('queued','retrying')
                      ) AS has_pending
               FROM batch_campaigns c
               WHERE c.status='in_progress' AND c.id>?
               ORDER BY c.id LIMIT ?""",
            (after_id, OVERDUE_CANCEL_CAMPAIGNS_PER_RUN),
        ).fetchall()

    rows = read_transaction(
        lambda conn: load_candidates(conn, _overdue_campaign_cursor),
        operation="batch_campaign.overdue_candidates",
    )
    if not rows and _overdue_campaign_cursor:
        _overdue_campaign_cursor = 0
        rows = read_transaction(
            lambda conn: load_candidates(conn, 0),
            operation="batch_campaign.overdue_candidates",
        )
    cancelled = 0
    for row in rows:
        _overdue_campaign_cursor = int(row["id"])
        start = (row["send_start_time"] or "")[:5]
        end = (row["send_end_time"] or "")[:5]
        if start and end:
            if start <= end:
                window_ended = current > end
            else:
                # Window crosses midnight: active while current >= start or current <= end
                window_ended = current > end and current < start
        elif end:
            window_ended = current > end
        else:
            continue
        if not window_ended:
            continue
        if not row["has_pending"]:
            continue

        def cancel(conn, campaign_id=row["id"]):
            task_ids = conn.execute(
                """SELECT id FROM sms_tasks
                   WHERE campaign_type='batch' AND campaign_id=?
                     AND status IN ('queued','retrying')
                   ORDER BY id LIMIT ?""",
                (campaign_id, OVERDUE_CANCEL_BATCH_SIZE),
            ).fetchall()
            ids = [int(task["id"]) for task in task_ids]
            if not ids:
                return []
            placeholders = ",".join("?" for _ in ids)
            sql = (
                "UPDATE sms_tasks SET status='cancelled', "
                "error_msg='daily send window ended', updated_at=CURRENT_TIMESTAMP "
                "WHERE campaign_type='batch' AND campaign_id=? "
                "AND status IN ('queued','retrying') AND id IN ("
                + placeholders + ") RETURNING id"
            )
            return conn.execute(
                sql,
                [campaign_id, *ids],
            ).fetchall()

        # One campaign per short transaction prevents a large campaign's
        # overdue cleanup from holding the writer gate while other campaigns
        # and the SMS result writer wait behind it.
        changed = write_transaction(
            cancel,
            attempts=3,
            operation="batch_campaign.cancel_overdue",
            base_delay=0.1,
        )
        cancelled += len(changed or [])
    if cancelled:
        logger.info("Batch overdue daily tasks cancelled=%s", cancelled)
    return cancelled


class _SequentialBatchSource:
    """A resumable sequential reader shared by one campaign's fair turns.

    The iterator is intentionally kept outside DuckDB.  A normal fair turn
    resumes the same reader in O(quantum) work; after a crash a new reader
    advances once to the durable manifest/chunk offset.  This avoids seeking
    from row zero for every batch while retaining a recoverable boundary.
    """

    def __init__(self, path: str, batch_size: int):
        self.path = path
        self.batch_size = max(1, int(batch_size or 500))
        self.ext = os.path.splitext(path)[1].lower()
        self.reader = None
        self.next_index = 0
        self.lookahead = None
        self.open()

    def open(self):
        self.close()
        self.reader = _iter_customer_list_frames(self.path, self.batch_size)
        self.next_index = 0
        self.lookahead = None

    def close(self):
        reader = self.reader
        self.reader = None
        if reader is not None:
            close = getattr(reader, "close", None)
            if close:
                try:
                    close()
                except Exception:
                    pass

    def frame_at(self, batch_index: int):
        batch_index = int(batch_index)
        if self.lookahead is not None and self.lookahead[0] == batch_index:
            frame = self.lookahead[1]
            self.lookahead = None
            return frame
        if self.lookahead is not None and self.lookahead[0] < batch_index:
            self.lookahead = None
        if batch_index < self.next_index:
            self.open()
        while self.next_index < batch_index:
            try:
                next(self.reader)
            except StopIteration:
                return None
            self.next_index += 1
        try:
            frame = next(self.reader)
        except StopIteration:
            return None
        self.next_index += 1
        return frame

    def peek_at(self, batch_index: int):
        frame = self.frame_at(batch_index)
        if frame is not None:
            self.lookahead = (int(batch_index), frame)
        return frame


def _source_signature(path: str) -> dict:
    stat = os.stat(path)
    return {
        "path": os.path.abspath(path),
        "size": int(stat.st_size),
        "mtime_ns": int(getattr(stat, "st_mtime_ns", int(stat.st_mtime * 1_000_000_000))),
    }


def _source_hash(path: str) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        while True:
            block = handle.read(1024 * 1024)
            if not block:
                break
            digest.update(block)
    return digest.hexdigest()


def _manifest_for_campaign(campaign_id: int, manifest_id: Optional[int] = None) -> Optional[dict]:
    def read(conn):
        if manifest_id is not None:
            row = conn.execute(
                "SELECT * FROM batch_campaign_manifests WHERE id=? AND campaign_id=?",
                (int(manifest_id), int(campaign_id)),
            ).fetchone()
        else:
            row = conn.execute(
                """SELECT * FROM batch_campaign_manifests
                   WHERE campaign_id=? AND status IN ('building','ready')
                   ORDER BY CASE WHEN status='building' THEN 0 ELSE 1 END,
                            updated_at DESC,id DESC LIMIT 1""",
                (int(campaign_id),),
            ).fetchone()
        return dict(row) if row else None
    return read_transaction(read, operation="batch_campaign.manifest.get")


def _manifest_matches(manifest: Optional[dict], signature: dict, batch_size: int) -> bool:
    return bool(
        manifest
        and os.path.abspath(str(manifest.get("source_path") or "")) == signature["path"]
        and int(manifest.get("source_size", 0) or 0) == signature["size"]
        and int(manifest.get("source_mtime_ns", 0) or 0) == signature["mtime_ns"]
        and int(manifest.get("batch_size", 0) or 0) == int(batch_size)
    )


def _create_manifest(campaign: dict, signature: dict, batch_size: int) -> dict:
    source_filename = campaign.get("customer_list_filename") or os.path.basename(signature["path"])

    def write(conn):
        # A stale builder must never become the active manifest. It remains in
        # the table for post-incident inspection and can be superseded safely.
        conn.execute(
            """UPDATE batch_campaign_manifests SET status='failed',
                      error_detail='superseded by newer source identity',updated_at=CURRENT_TIMESTAMP
               WHERE campaign_id=? AND status='building'""",
            (int(campaign["id"]),),
        )
        row = conn.execute(
            """INSERT INTO batch_campaign_manifests
                 (campaign_id,source_path,source_filename,source_size,source_mtime_ns,
                  batch_size,status)
               VALUES (?,?,?,?,?,?, 'building') RETURNING *""",
            (int(campaign["id"]), signature["path"], source_filename,
             signature["size"], signature["mtime_ns"], int(batch_size)),
        ).fetchone()
        return dict(row)
    return _run_campaign_write("manifest.create", write)


def _manifest_batches(manifest_id: int) -> list[dict]:
    rows = read_transaction(
        lambda conn: conn.execute(
            "SELECT * FROM batch_campaign_batches WHERE manifest_id=? ORDER BY batch_index",
            (int(manifest_id),),
        ).fetchall(),
        operation="batch_campaign.manifest_batches",
    )
    return [dict(row) for row in rows]


def _ensure_batch_ingestion(campaign: dict, manifest: dict) -> None:
    """Publish the chunk lease manifest idempotently after generation commits."""
    batches = _manifest_batches(manifest["id"])
    existing = ingestion_job_service.get_file(
        "batch", campaign["id"],
        campaign.get("customer_list_filename") or os.path.basename(manifest["source_path"]),
    )
    same_manifest = existing and int(existing.get("manifest_id") or 0) == int(manifest["id"])
    if same_manifest and int(existing.get("total_chunks") or 0) == len(batches):
        return
    ingestion_job_service.register_file(
        "batch", campaign["id"],
        campaign.get("customer_list_filename") or os.path.basename(manifest["source_path"]),
        manifest["source_path"], total_rows=int(manifest.get("total_rows") or 0),
        chunk_size=int(manifest.get("batch_size") or 500),
        chunks=[
            {"source_batch_id": item["id"], "chunk_index": item["batch_index"],
             "start_row": item["start_row"], "end_row": item["end_row"]}
            for item in batches
        ], reset=bool(existing and not same_manifest), manifest_id=int(manifest["id"]),
    )


def _publish_manifest(campaign_id: int, manifest_id: int, total_rows: int,
                      source_hash: str = "") -> dict:
    def write(conn):
        # Hashing happens after the source has been fully read, outside this
        # short publication transaction.  The hash itself is published in the
        # same visibility flip as the ready state.
        conn.execute(
            """UPDATE batch_campaign_manifests SET status='ready',total_rows=?,
                      total_batches=generated_batches,source_hash=?,published_at=CURRENT_TIMESTAMP,
                      updated_at=CURRENT_TIMESTAMP
               WHERE id=? AND campaign_id=? AND status='building'""",
            (int(total_rows), source_hash, int(manifest_id), int(campaign_id)),
        )
        conn.execute(
            """UPDATE batch_campaigns SET active_manifest_id=?,updated_at=CURRENT_TIMESTAMP
               WHERE id=?""",
            (int(manifest_id), int(campaign_id)),
        )
        return dict(conn.execute(
            "SELECT * FROM batch_campaign_manifests WHERE id=?", (int(manifest_id),)
        ).fetchone())
    return _run_campaign_write("manifest.publish", write)


def _mark_manifest_failed(manifest_id: int, error: Exception) -> None:
    try:
        _run_campaign_write(
            "manifest.fail",
            lambda conn: conn.execute(
                """UPDATE batch_campaign_manifests SET status='failed',error_detail=?,
                          updated_at=CURRENT_TIMESTAMP WHERE id=? AND status='building'""",
                (str(error)[:4000], int(manifest_id)),
            ),
        )
    except Exception:
        logger.exception("Unable to mark batch manifest failed manifest=%s", manifest_id)


def pending_campaign_ids() -> list[int]:
    rows = read_transaction(
        lambda conn: conn.execute(
            """SELECT id FROM batch_campaigns
               WHERE status='in_progress' AND COALESCE(customer_list_path,'')<>''
               ORDER BY id"""
        ).fetchall(),
        operation="batch_campaign.pending_campaigns",
    )
    return [int(row["id"]) for row in rows]


def _manifest_ready_for_campaign(campaign: dict) -> Optional[dict]:
    manifest = _manifest_for_campaign(campaign["id"], campaign.get("active_manifest_id"))
    return manifest if manifest and manifest.get("status") == "ready" else None


def generate_batches(campaign_id: int, batch_size: int = 500,
                     max_batches: Optional[int] = None) -> dict:
    """Build a durable, invisible manifest in bounded write transactions.

    ``max_batches`` is used by the scheduler as its fairness quantum. API
    callers retain the historical full-generation behavior when it is omitted.
    """
    campaign = get_one(campaign_id)
    pressure, _ = backpressure_service.reasons(producing_chunks=True)
    if pressure:
        return {"success": False, "message": "ingestion backpressure: " + ",".join(pressure)}
    if not campaign:
        return {"success": False, "message": "Campaign not found"}
    path = campaign.get("customer_list_path", "")
    if not path or not os.path.isfile(path):
        return {"success": False, "message": "No customer list uploaded"}
    batch_size = max(1, int(batch_size or 500))
    try:
        signature = _source_signature(path)
    except OSError as exc:
        return {"success": False, "message": f"Cannot read file: {exc}"}

    with _batch_generation_lock:
        manifest = _manifest_for_campaign(campaign_id)
        if not _manifest_matches(manifest, signature, batch_size):
            manifest = _create_manifest(campaign, signature, batch_size)
        if manifest.get("status") == "ready":
            _ensure_batch_ingestion(campaign, manifest)
            batches = _manifest_batches(manifest["id"])
            return {"success": True, "batches": [
                {**item, "count": item.get("total_count", 0)} for item in batches
            ], "total_batches": len(batches), "total_rows": manifest.get("total_rows", 0),
                    "manifest_id": manifest["id"], "complete": True}

        key = (int(manifest["id"]), signature["path"], batch_size)
        with _batch_source_lock:
            source = _batch_generation_cursors.get(key)
            if source is None:
                source = _SequentialBatchSource(path, batch_size)
                _batch_generation_cursors[key] = source
            generated = int(manifest.get("generated_batches", 0) or 0)
            total_rows = int(manifest.get("total_rows", 0) or 0)
            # A process restart has no cursor; advance exactly once to the
            # durable generated offset. Normal fair turns continue in place.
            if source.next_index > generated:
                source.open()
            local_created = []
            limit = None if max_batches is None else max(1, int(max_batches))
            try:
                while limit is None or len(local_created) < limit:
                    frame = source.frame_at(generated)
                    if frame is None:
                        break
                    count = len(frame)
                    start = total_rows
                    total_rows += count

                    def write_batch(conn, index=generated, start_row=start,
                                    end_row=total_rows, item_count=count):
                        existing = conn.execute(
                            "SELECT * FROM batch_campaign_batches WHERE manifest_id=? AND batch_index=?",
                            (int(manifest["id"]), int(index)),
                        ).fetchone()
                        if existing:
                            return dict(existing)
                        row = conn.execute(
                            """INSERT INTO batch_campaign_batches
                               (campaign_id,manifest_id,batch_index,start_row,end_row,total_count)
                               VALUES (?,?,?,?,?,?) RETURNING *""",
                            (int(campaign_id), int(manifest["id"]), int(index),
                             int(start_row), int(end_row), int(item_count)),
                        ).fetchone()
                        conn.execute(
                            """UPDATE batch_campaign_manifests SET generated_batches=?,
                                      total_rows=?,updated_at=CURRENT_TIMESTAMP WHERE id=?""",
                            (int(index + 1), int(end_row), int(manifest["id"])),
                        )
                        return dict(row)

                    batch = _run_campaign_write("manifest.batch", write_batch)
                    local_created.append({**batch, "count": batch.get("total_count", count)})
                    generated += 1
                    del frame
                source_exhausted = source.peek_at(generated) is None
                if source_exhausted:
                    # The probe above advances no state beyond EOF and keeps
                    # the source cursor reusable for diagnostics.
                    final_hash = _source_hash(path)
                    current_signature = _source_signature(path)
                    if current_signature != signature:
                        raise RuntimeError("customer list changed during manifest generation")
                    manifest = _publish_manifest(
                        campaign_id, manifest["id"], total_rows, source_hash=final_hash,
                    )
                    _ensure_batch_ingestion(campaign, manifest)
                    _batch_generation_cursors.pop(key, None)
                    return {"success": True, "batches": [
                        {**item, "count": item.get("total_count", 0)}
                        for item in _manifest_batches(manifest["id"])
                    ], "total_batches": generated, "total_rows": total_rows,
                            "manifest_id": manifest["id"], "complete": True,
                            "source_hash": final_hash}
                return {"success": True, "batches": local_created,
                        "total_batches": generated, "total_rows": total_rows,
                        "manifest_id": manifest["id"], "complete": False}
            except Exception as exc:
                _batch_generation_cursors.pop(key, None)
                _mark_manifest_failed(manifest["id"], exc)
                return {"success": False, "message": str(exc), "manifest_id": manifest["id"]}


def get_batches(campaign_id: int) -> list:
    rows = read_transaction(
        lambda conn: conn.execute(
            """SELECT b.* FROM batch_campaign_batches b
               JOIN batch_campaigns c ON c.id=b.campaign_id
               WHERE b.campaign_id=?
                 AND (c.active_manifest_id IS NULL OR b.manifest_id=c.active_manifest_id)
               ORDER BY b.batch_index""",
            (campaign_id,)
        ).fetchall(),
        operation="batch_campaign.batches",
    )
    return [dict(r) for r in rows]


def _get_today_sent_msisdns(campaign_dir: str) -> set:
    """Read today's output file and return set of already-sent msisdns."""
    today_str = datetime.now().strftime("%Y%m%d")
    output_path = os.path.join(campaign_dir, "output", f"{today_str}.csv")
    if not os.path.isfile(output_path):
        return set()
    try:
        df = pd.read_csv(output_path, dtype=str, encoding="utf-8-sig")
        if "msisdn" in df.columns:
            return set(df["msisdn"].dropna().str.strip())
    except Exception:
        pass
    return set()


def process_batch(batch_id: int, batch_rows=None, campaign_override: Optional[dict] = None) -> dict:
    def read_batch(conn):
        row = conn.execute(
            "SELECT * FROM batch_campaign_batches WHERE id=?", (batch_id,)
        ).fetchone()
        if not row:
            return None, None, None
        active_manifest = conn.execute(
            "SELECT active_manifest_id FROM batch_campaigns WHERE id=?",
            (row["campaign_id"],),
        ).fetchone()
        return dict(row), row["campaign_id"], (
            active_manifest["active_manifest_id"] if active_manifest else None
        )

    batch, campaign_id, active_manifest_id = read_transaction(
        read_batch,
        operation="batch.process.snapshot",
    )
    if not batch:
        return {"success": False, "message": "Batch not found"}
    if active_manifest_id is not None and int(batch.get("manifest_id") or 0) != int(active_manifest_id):
        return {"success": False, "message": "Batch belongs to an inactive manifest"}

    if batch["status"] not in {"pending", "failed"}:
        return {"success": False, "message": "Batch is already processing or queued"}

    if batch_rows is None:
        campaign_probe = get_one(campaign_id)
        if campaign_probe and campaign_probe.get("customer_list_path"):
            return process_pending_batches(campaign_id, batch_id)

    chunk_job = ingestion_job_service.claim_chunk(source_batch_id=batch_id)
    if chunk_job and chunk_job.get("status") == "completed":
        return {"success": True, "message": "Chunk already completed", "queued_count": 0,
                "resumed": True}
    if chunk_job is None:
        mapped = read_transaction(
            lambda conn: conn.execute(
                "SELECT id FROM ingestion_chunks WHERE source_batch_id=?", (batch_id,)
            ).fetchone(),
            operation="batch.process.chunk_lease_check",
        )
        if mapped:
            return {"success": False, "message": "Chunk lease is held by another worker"}

    write_transaction(
        lambda conn: conn.execute(
            "UPDATE batch_campaign_batches SET status='processing', updated_at=datetime('now','localtime') WHERE id=?",
            (batch_id,),
        ),
        attempts=3,
        operation="batch.process.mark_processing",
    )
    try:
        campaign = campaign_override or get_one(campaign_id)
        path = campaign.get("customer_list_path", "") if campaign else ""
        if not path or not os.path.isfile(path):
            raise Exception("Customer list file not found")
        if batch_rows is None:
            raise RuntimeError("batch rows must be supplied by the sequential ingestion reader")

        batch_rows = _canonicalize_customer_list_columns(batch_rows)
        if "msisdn" not in batch_rows.columns:
            raise ValueError("Customer list file must contain an msisdn column")
        enqueue_result = sms_task_service.enqueue_rows(
            "batch", campaign, batch_rows.fillna("").to_dict("records"),
            campaign.get("customer_list_filename") or os.path.basename(path), batch_id
        )
        if not enqueue_result.get("success"):
            raise Exception(enqueue_result.get("message", "Unable to enqueue batch"))
        queued = enqueue_result["created"]
        batch_status = "queued" if queued > 0 else "skipped"

        def complete(conn):
            conn.execute(
                """UPDATE batch_campaign_batches
                   SET status=?,task_count=?,delivered_count=0,failed_count=0,
                       updated_at=datetime('now','localtime') WHERE id=?""",
                (batch_status, queued, batch_id),
            )
            # Update daily stats in the same short commit as the batch state.
            stats_date = datetime.now().strftime("%Y-%m-%d")
            total_prospects = campaign.get("row_count", 0) if campaign else 0
            existing_stats = conn.execute(
                "SELECT id, sent_count FROM batch_campaign_daily_stats WHERE campaign_id=? AND stats_date=?",
                (campaign_id, stats_date),
            ).fetchone()
            if existing_stats:
                conn.execute(
                    "UPDATE batch_campaign_daily_stats SET updated_at=datetime('now','localtime') WHERE id=?",
                    (existing_stats["id"],),
                )
            else:
                conn.execute(
                    "INSERT INTO batch_campaign_daily_stats (campaign_id, stats_date, total_prospects, sent_count) VALUES (?, ?, ?, ?)",
                    (campaign_id, stats_date, total_prospects, 0),
                )
            if chunk_job:
                ingestion_job_service.complete_chunk_in_transaction(
                    conn, chunk_job["id"], enqueue_result,
                )

        write_transaction(
            complete,
            attempts=3,
            operation="batch.process.complete",
        )
        return {"success": True, "message": f"Queued {queued} messages",
                "queued_count": queued, "filter_result": enqueue_result}
    except Exception as e:
        if is_transient_write_error(e):
            raise
        def fail(conn):
            conn.execute(
                "UPDATE batch_campaign_batches SET status='failed', updated_at=datetime('now','localtime') WHERE id=?",
                (batch_id,),
            )
            if chunk_job:
                ingestion_job_service.fail_chunk_in_transaction(conn, chunk_job["id"], e)

        write_transaction(
            fail,
            attempts=3,
            operation="batch.process.fail",
        )
        return {"success": False, "message": str(e)}


def process_pending_batches(campaign_id: Optional[int] = None, target_batch_id: Optional[int] = None,
                            stop_event=None, max_batches: Optional[int] = None,
                            quantum_seconds: Optional[float] = None) -> dict:
    """Process a bounded batch quantum and yield the shared worker.

    The per-manifest source cursor advances sequentially and is keyed by the
    durable manifest identity.  A restart reconstructs it once to the first
    incomplete batch; ordinary scheduler turns never rescan the file prefix.
    """
    if campaign_id is None:
        pending = pending_campaign_ids()
        campaign_id = pending[0] if pending else None
    if campaign_id is None:
        return {"success": True, "processed": 0}

    def read_pending(conn):
        campaign_row = conn.execute(
            "SELECT * FROM batch_campaigns WHERE id=? AND status='in_progress'", (campaign_id,)
        ).fetchone()
        if campaign_row:
            active_manifest_id = campaign_row.get("active_manifest_id")
            batches = conn.execute(
                """SELECT * FROM batch_campaign_batches
                   WHERE campaign_id=? AND
                     ((? IS NULL AND manifest_id IS NULL) OR manifest_id=?)
                   ORDER BY batch_index""",
                (campaign_id, active_manifest_id, active_manifest_id),
            ).fetchall()
        else:
            batches = []
        return campaign_row, batches

    campaign_row, batches = read_transaction(
        read_pending,
        operation="batch_campaign.pending_snapshot",
    )
    campaign = _enrich(dict(campaign_row)) if campaign_row else None
    if not campaign:
        return {"success": False, "message": "Campaign not found", "processed": 0}
    manifest = _manifest_for_campaign(campaign_id, campaign.get("active_manifest_id"))
    if manifest and manifest.get("status") != "ready":
        return {"success": True, "message": "batch manifest is still building", "processed": 0}
    if not batches:
        return {"success": True, "message": "no visible batches", "processed": 0}
    if manifest:
        _ensure_batch_ingestion(campaign, manifest)
    active, reason = sms_task_service.campaign_is_active(campaign)
    if not active:
        return {"success": False, "message": reason, "processed": 0}
    path = campaign.get("customer_list_path", "")
    if not path or not os.path.isfile(path):
        return {"success": False, "message": "Customer list file not found", "processed": 0}

    batch_dicts = [dict(row) for row in batches]
    target = next((row for row in batch_dicts if target_batch_id is not None and
                   int(row["id"]) == int(target_batch_id)), None)
    if target_batch_id is not None and target is None:
        return {"success": False, "message": "Batch not found", "processed": 0}
    pending = [row for row in batch_dicts if row.get("status") in {"pending", "failed"}]
    if target is not None:
        pending = [target]
    if not pending:
        return {"success": True, "processed": 0, "complete": True}

    batch_index = int(pending[0].get("batch_index") or 0)
    manifest_key = int((manifest or {}).get("id") or 0)
    try:
        signature = _source_signature(path)
    except OSError as exc:
        return {"success": False, "message": str(exc), "processed": 0}
    source_key = (int(campaign_id), manifest_key, signature["path"],
                  signature["size"], signature["mtime_ns"],
                  int(campaign.get("batch_size") or 500))
    processed = 0
    started = time.monotonic()
    limit = None if max_batches is None else max(1, int(max_batches))
    with _batch_source_lock:
        source = _batch_source_cursors.get(source_key)
        if source is None:
            source = _SequentialBatchSource(path, int(campaign.get("batch_size") or 500))
            _batch_source_cursors[source_key] = source
        try:
            for batch in pending:
                if limit is not None and processed >= limit:
                    break
                if quantum_seconds is not None and processed and time.monotonic() - started >= quantum_seconds:
                    break
                if stop_event is not None and stop_event.is_set():
                    break
                active, reason = sms_task_service.campaign_is_active(campaign)
                if not active:
                    return {"success": False, "message": reason, "processed": processed}
                if not _wait_for_ingest_capacity(stop_event):
                    return {"success": False, "message": "batch ingestion worker stopping", "processed": processed}
                frame = source.frame_at(int(batch.get("batch_index") or 0))
                if frame is None:
                    raise RuntimeError(f"batch source ended before batch {batch.get('batch_index')}")
                result = process_batch(batch["id"], batch_rows=frame, campaign_override=campaign)
                del frame
                if result.get("success"):
                    processed += 1
                else:
                    # Do not let a failed frame advance the in-memory cursor;
                    # its durable chunk remains retryable on the next turn.
                    source.open()
                    break
            if processed:
                metrics_service.observe_histogram(
                    "marketing_copilot_batch_quantum_seconds",
                    time.monotonic() - started,
                    labels={"campaign_id": campaign_id},
                )
        except Exception as exc:
            source.open()
            logger.exception("Sequential batch read failed campaign=%s", campaign_id)
            return {"success": False, "message": str(exc), "processed": processed}

    manifest_file = ingestion_job_service.get_file(
        "batch", campaign_id,
        campaign.get("customer_list_filename") or os.path.basename(path),
    )
    if manifest_file and processed:
        ingestion_job_service.finish_file(manifest_file["id"])
    return {"success": True, "processed": processed,
            "complete": processed >= len(pending)}


def _count_reached_users(campaign_dir: str) -> int:
    output_dir = os.path.join(campaign_dir, "output")
    if not os.path.isdir(output_dir):
        return 0
    unique = set()
    for fname in os.listdir(output_dir):
        if not fname.endswith(".csv"):
            continue
        path = os.path.join(output_dir, fname)
        try:
            df = pd.read_csv(path, dtype=str, encoding="utf-8-sig", usecols=["msisdn"])
            unique.update(df["msisdn"].dropna().str.strip().tolist())
        except Exception:
            pass
    return len(unique)


def get_detail(campaign_id: int) -> Optional[dict]:
    campaign = get_one(campaign_id)
    if not campaign:
        return None
    campaign["daily_stats"] = get_daily_stats(campaign_id)
    _attach_daily_smpp(campaign_id, campaign["daily_stats"])
    campaign["batches"] = get_batches(campaign_id)
    reached = read_transaction(
        lambda conn: conn.execute(
            """SELECT COUNT(DISTINCT msisdn) cnt FROM sms_tasks
               WHERE campaign_type='batch' AND campaign_id=? AND status IN ('delivered','accepted')""",
            (campaign_id,),
        ).fetchone(),
        operation="batch_campaign.detail.reached",
    )
    campaign["reached_users"] = reached["cnt"] if reached else 0
    campaign["sms_templates"] = _get_campaign_templates(campaign_id)
    return campaign


def _attach_daily_smpp(campaign_id: int, daily_stats: list) -> None:
    """Attach the SMSC label(s) used per stats date, derived from sms_tasks."""
    if not daily_stats:
        return
    rows, config_rows = read_transaction(
        lambda conn: (
            conn.execute(
                """SELECT CAST(created_at AS DATE) AS d, smpp_config_id
                   FROM sms_tasks
                   WHERE campaign_type='batch' AND campaign_id=?
                   GROUP BY d, smpp_config_id""",
                (campaign_id,),
            ).fetchall(),
            conn.execute("SELECT id, name, host, port FROM smpp_configs").fetchall(),
        ),
        operation="batch_campaign.detail.smpp",
    )
    configs = {r["id"]: r for r in config_rows}

    def label(config_id):
        if config_id is None:
            return "默认"
        cfg = configs.get(config_id)
        if cfg:
            return f"{cfg['name']} ({cfg['host']}:{cfg['port']})"
        return f"#{config_id}"

    by_date = {}
    for r in rows:
        by_date.setdefault(str(r["d"]), []).append(r["smpp_config_id"])
    for item in daily_stats:
        ids = by_date.get(item["stats_date"], [])
        item["smpp_label"] = "、".join(dict.fromkeys(label(i) for i in ids)) or "-"


def _publish_customer_list_atomic(source_path: str, destination_path: str) -> None:
    """Publish a validated target list without exposing a partial file."""
    try:
        os.replace(source_path, destination_path)
        return
    except OSError as exc:
        if exc.errno != errno.EXDEV:
            raise

    staged_path = f"{destination_path}.upload-{uuid.uuid4().hex}.tmp"
    try:
        with open(source_path, "rb") as source, open(staged_path, "wb") as staged:
            shutil.copyfileobj(source, staged, length=1024 * 1024)
        os.replace(staged_path, destination_path)
    finally:
        try:
            if os.path.exists(staged_path):
                os.unlink(staged_path)
        except OSError:
            logger.warning("Unable to remove staged target-list file=%s", staged_path, exc_info=True)


def upload_customer_list(campaign_id: int, file_path: str, original_filename: str) -> Optional[dict]:
    campaign = get_one(campaign_id)
    if not campaign:
        return None

    campaign_dir = campaign.get("campaign_dir", "")
    if not campaign_dir:
        campaign_dir = _ensure_campaign_dir(campaign)
        os.makedirs(os.path.join(campaign_dir, "whitelist"), exist_ok=True)
        os.makedirs(os.path.join(campaign_dir, "output"), exist_ok=True)
        write_transaction(
            lambda conn: conn.execute(
                "UPDATE batch_campaigns SET campaign_dir=? WHERE id=?",
                (campaign_dir, campaign_id),
            ),
            operation="batch_campaign.directory",
        )

    whitelist_dir = os.path.join(campaign_dir, "whitelist")
    os.makedirs(whitelist_dir, exist_ok=True)
    original_filename = os.path.basename(original_filename or "target-list.csv")
    ext = os.path.splitext(original_filename)[1].lower()
    if ext not in BATCH_CUSTOMER_LIST_EXTENSIONS:
        raise ValueError("Only CSV, TXT, XLSX and XLS files are supported")

    try:
        validation = _validate_customer_list_file(file_path)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"Cannot read customer list file: {exc}") from exc

    dest = os.path.join(whitelist_dir, f"whitelist{ext}")
    _publish_customer_list_atomic(file_path, dest)
    row_count = int(validation["row_count"])
    ignored_rows = int(validation.get("ignored_rows", 0) or 0)
    if ignored_rows:
        logger.warning(
            "Batch target list ignored rows campaign=%s total=%s invalid_msisdn=%s malformed=%s",
            campaign_id, ignored_rows, validation.get("ignored_msisdn_rows", 0),
            validation.get("ignored_malformed_rows", 0),
        )

    def write(conn):
        conn.execute(
            "UPDATE batch_campaigns SET customer_list_filename=?, customer_list_path=?, row_count=?, updated_at=datetime('now','localtime') WHERE id=?",
            (original_filename, dest, row_count, campaign_id),
        )
        return conn.execute(
            "SELECT * FROM batch_campaigns WHERE id=?", (campaign_id,)
        ).fetchone()

    row = write_transaction(write, operation="batch_campaign.upload")
    sms_task_service.invalidate_campaign_cache("batch", campaign_id)

    # If campaign is in progress, regenerate batches from new list (no auto-send)
    if campaign.get("status") == "in_progress" and row_count > 0:
        batch_size = campaign.get("batch_size") or 500
        generate_batches(campaign_id, batch_size)

    if not row:
        return None
    return {
        **dict(row),
        "ignored_rows": ignored_rows,
        "ignored_msisdn_rows": int(validation.get("ignored_msisdn_rows", 0) or 0),
        "ignored_malformed_rows": int(validation.get("ignored_malformed_rows", 0) or 0),
    }
